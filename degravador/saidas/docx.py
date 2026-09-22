"""Saída DOCX — documento oficial de degravação.

Segue ABNT NBR 14724 + Manual de Redação da Presidência da República:
Arial 12, justificado, A4, margens 3/3/2/2 cm, entrelinha 1,5, recuo de 1,25 cm,
paginação no canto superior direito. Os parâmetros ficam em
``estilos/oficial_abnt.json`` (instituições customizam sem tocar no código).

A FALA é reproduzida fielmente; apenas os textos redigidos pela aplicação
(cabeçalho, legendas, termo de encerramento) seguem a norma culta.
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from ._comum import iter_turnos, fmt_hms, rotulo_idioma, rotulo_locutor


_MESES = [
    "janeiro", "fevereiro", "março", "abril", "maio", "junho",
    "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
]

# Estilo padrão embarcado; sobrescrito pelo arquivo estilos/oficial_abnt.json.
_ESTILO_PADRAO = {
    "titulo": "TERMO DE DEGRAVAÇÃO",
    "fonte": "Arial",
    "tamanho_pt": 12,
    "cor_hex": "000000",
    "espacamento_linhas": 1.5,
    "recuo_primeira_linha_cm": 1.25,
    "margens_cm": {"superior": 3, "esquerda": 3, "inferior": 2, "direita": 2},
    "cor_realce_baixa_confianca_hex": "F2DEDE",
    "termo_encerramento": (
        "Documento produzido por processamento automatizado local (VOX)."
    ),
}


def _carregar_estilo(caminho_estilo: str | Path | None) -> dict:
    estilo = dict(_ESTILO_PADRAO)
    if caminho_estilo and Path(caminho_estilo).exists():
        with open(caminho_estilo, "r", encoding="utf-8") as f:
            estilo.update(json.load(f))
    return estilo


def _data_extenso(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        dt = datetime.fromisoformat(iso)
    except ValueError:
        return iso
    return f"{dt.day} de {_MESES[dt.month - 1]} de {dt.year}"


def _sombrear_run(run, cor_hex: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rpr = run._element.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor_hex)
    rpr.append(shd)


def _campo_pagina(paragrafo) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    run._r.append(inicio)
    run._r.append(instr)
    run._r.append(fim)


def _locutores_em_ordem(doc: dict) -> list[str]:
    vistos: list[str] = []
    for seg in doc.get("segments", []):
        spk = seg.get("speaker")
        if spk and spk not in vistos:
            vistos.append(spk)
    return vistos


def exportar(
    doc: dict,
    caminho: str | Path,
    *,
    mapa_locutores: dict | None = None,
    sem_timestamps: bool = False,
    caminho_estilo: str | Path | None = None,
    **_opts,
) -> Path:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION

    estilo = _carregar_estilo(caminho_estilo)
    meta = doc.get("metadata", {})
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    documento = Document()

    # --- Estilo base (Normal): Arial 12, justificado, 1,5, recuo -----------
    normal = documento.styles["Normal"]
    normal.font.name = estilo["fonte"]
    normal.font.size = Pt(estilo["tamanho_pt"])
    cor = estilo["cor_hex"]
    normal.font.color.rgb = RGBColor(int(cor[0:2], 16), int(cor[2:4], 16), int(cor[4:6], 16))
    pf = normal.paragraph_format
    pf.line_spacing = estilo["espacamento_linhas"]
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # --- Página A4 + margens ABNT ------------------------------------------
    sec = documento.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    m = estilo["margens_cm"]
    sec.top_margin = Cm(m["superior"])
    sec.left_margin = Cm(m["esquerda"])
    sec.bottom_margin = Cm(m["inferior"])
    sec.right_margin = Cm(m["direita"])

    # --- Paginação no canto superior direito -------------------------------
    cab = sec.header.paragraphs[0]
    cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _campo_pagina(cab)

    # === 1. Cabeçalho de identificação =====================================
    titulo = documento.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(estilo["titulo"].upper())
    run.bold = True

    def _linha_meta(rotulo: str, valor) -> None:
        p = documento.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.add_run(f"{rotulo}: ").bold = True
        p.add_run(str(valor if valor is not None else "—"))

    documento.add_paragraph()
    _linha_meta("Arquivo de origem", meta.get("arquivo"))
    _linha_meta("Hash SHA-256 do original", meta.get("sha256"))
    _linha_meta("Data e hora do processamento", meta.get("criado_em"))
    dur = meta.get("duracao_s")
    _linha_meta("Duração do áudio", fmt_hms(dur) if dur else None)
    _linha_meta("Interlocutores identificados", meta.get("num_speakers"))
    _linha_meta("Idioma da fala", rotulo_idioma(meta.get("params")))
    _linha_meta("Modelo utilizado", meta.get("modelo"))
    _linha_meta("Versão do Degravador", meta.get("versao_app"))
    _linha_meta("Perfil de processamento", meta.get("perfil"))
    _linha_meta("Responsável pela conferência", meta.get("revisado_por") or "____________________")

    # === 2. Legenda de interlocutores ======================================
    locutores = _locutores_em_ordem(doc)
    if locutores:
        documento.add_paragraph()
        cab2 = documento.add_paragraph()
        cab2.add_run("Legenda de interlocutores").bold = True
        renomeados = bool(mapa_locutores)
        for spk in locutores:
            nome = rotulo_locutor(spk, mapa_locutores)
            p = documento.add_paragraph()
            p.paragraph_format.line_spacing = 1.0
            if nome and nome != spk:
                p.add_run(f"{spk} → {nome}")
            else:
                p.add_run(f"{spk} (rótulo técnico — não renomeado)")
        if not renomeados:
            nota = documento.add_paragraph()
            nota.paragraph_format.line_spacing = 1.0
            nota.add_run(
                "Nota: os rótulos técnicos foram mantidos por não terem sido "
                "atribuídos nomes pelo responsável."
            ).italic = True

    # === 3. Corpo da degravação ============================================
    documento.add_paragraph()
    corpo_titulo = documento.add_paragraph()
    corpo_titulo.add_run("Degravação").bold = True

    identificar = len(locutores) > 1
    cor_realce = estilo["cor_realce_baixa_confianca_hex"]

    for turno in iter_turnos(doc):
        p = documento.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(estilo["recuo_primeira_linha_cm"])

        if identificar:
            nome = (rotulo_locutor(turno["speaker"], mapa_locutores) or "LOCUTOR").upper()
            ts = "" if sem_timestamps else f" ({fmt_hms(turno['start'])})"
            p.add_run(f"{nome}{ts}: ").bold = True

        for seg in turno["segmentos"]:
            flags = set(seg.get("flags", []))
            if "inaudivel" in flags:
                p.add_run(f"[inaudível – {fmt_hms(seg.get('start', 0.0))}] ")
                continue
            palavras = seg.get("words", [])
            if palavras:
                for w in palavras:
                    texto = (w.get("word") or "").strip()
                    if not texto:
                        continue
                    run = p.add_run(texto + " ")
                    if "baixa_confianca" in set(w.get("flags", [])):
                        _sombrear_run(run, cor_realce)
            else:
                p.add_run((seg.get("text") or "").strip() + " ")
            if "vozes_sobrepostas" in flags:
                p.add_run("[vozes sobrepostas] ")

    # === 4. Termo de encerramento ==========================================
    # Sem alinhamento por palavra não há score por palavra: a ausência de
    # realces de baixa confiança significa "não medido", não "conferido".
    if (meta.get("params") or {}).get("alinhamento_por_palavra") is False:
        documento.add_paragraph()
        obs = documento.add_paragraph()
        run_obs = obs.add_run(
            "Observação: não houve alinhamento por palavra para este idioma. "
            "Os timestamps são por trecho de fala, e a confiança foi medida por "
            "trecho — não palavra a palavra. A ausência de realces de baixa "
            "confiança palavra a palavra, portanto, NÃO indica transcrição "
            "conferida: indica medição não realizada nesse nível."
        )
        run_obs.bold = True

    documento.add_paragraph()
    termo = documento.add_paragraph()
    termo.add_run(estilo["termo_encerramento"])
    data = _data_extenso(meta.get("criado_em"))
    if data:
        pdata = documento.add_paragraph()
        pdata.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pdata.add_run(data)
    documento.add_paragraph()
    assinatura = documento.add_paragraph()
    assinatura.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assinatura.add_run("____________________________________\nResponsável pela conferência")

    documento.save(str(caminho))
    return caminho
